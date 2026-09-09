from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"C:\Users\CHOIIH\Downloads\(서식1) 일본 온라인 플랫폼(큐텐) 입점지원 신청서_변환.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="000000", size="6", edges=("top", "left", "bottom", "right")):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in edges:
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=40, start=55, bottom=40, end=55):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_inches):
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(round(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_row_height(row, height_inches):
    tr_pr = row._tr.get_or_add_trPr()
    tr_height = OxmlElement("w:trHeight")
    tr_height.set(qn("w:val"), str(round(height_inches * 1440)))
    tr_height.set(qn("w:hRule"), "exact")
    tr_pr.append(tr_height)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, size=7.3, bold=False, color="000000", name="Malgun Gothic"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def clear_cell(cell):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    return p


def fill_cell(cell, runs, *, align=WD_ALIGN_PARAGRAPH.CENTER, valign=WD_CELL_VERTICAL_ALIGNMENT.CENTER,
              size=7.3, bold=False, color="000000"):
    p = clear_cell(cell)
    p.alignment = align
    cell.vertical_alignment = valign
    if isinstance(runs, str):
        runs = [(runs, {})]
    for text, style in runs:
        r = p.add_run(text)
        set_run_font(
            r,
            size=style.get("size", size),
            bold=style.get("bold", bold),
            color=style.get("color", color),
            name=style.get("name", "Malgun Gothic"),
        )
    return cell


def merge_rect(table, r0, c0, r1, c1):
    return table.cell(r0, c0).merge(table.cell(r1, c1))


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(59528 / 7200)
    section.page_height = Inches(84186 / 7200)
    section.left_margin = Inches(5669 / 7200)
    section.right_margin = Inches(5669 / 7200)
    section.top_margin = Inches(3600 / 7200)
    section.bottom_margin = Inches(3600 / 7200)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    normal = doc.styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(7.3)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    # Title band
    title = doc.add_table(rows=1, cols=2)
    title.alignment = WD_TABLE_ALIGNMENT.CENTER
    title.autofit = False
    set_cell_width(title.cell(0, 0), 0.94)
    set_cell_width(title.cell(0, 1), 5.70)
    set_cell_shading(title.cell(0, 0), "0047A6")
    set_cell_border(title.cell(0, 0), size="10", edges=("bottom",))
    set_cell_border(title.cell(0, 1), size="10", edges=("bottom",))
    set_cell_margins(title.cell(0, 0), top=55, bottom=55)
    set_cell_margins(title.cell(0, 1), top=55, start=170, bottom=55)
    fill_cell(title.cell(0, 0), "서식", size=15, bold=True, color="FFFFFF")
    fill_cell(
        title.cell(0, 1),
        "일본 온라인 플랫폼(큐텐) 입점지원 신청서",
        align=WD_ALIGN_PARAGRAPH.LEFT,
        size=14,
        bold=True,
    )

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(5)
    spacer.add_run(" ")

    table = doc.add_table(rows=21, cols=9)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    grid_units = [2552, 1481, 4033, 5921, 10045, 2114, 5356, 8123, 8123]
    grid_inches = [value / 7200 for value in grid_units]
    for row in table.rows:
        for idx, width in enumerate(grid_inches):
            set_cell_width(row.cells[idx], width)
        set_cant_split(row)
        for cell in row.cells:
            set_cell_border(cell, size="5")
            set_cell_margins(cell)

    heights = [2228, 1280, 1380, 1380, 1380, 1380, 1380, 1380, 1380, 2424, 2424,
               1663, 2707, 2707, 8086, 6571, 1663, 3556, 2990, 8836, 6698]
    for row, height in zip(table.rows, heights):
        set_row_height(row, height / 7200)

    # Company and trademark information
    fill_cell(merge_rect(table, 0, 0, 1, 2), "기 업 명", size=8.2)
    fill_cell(merge_rect(table, 0, 3, 1, 4), "")
    fill_cell(merge_rect(table, 0, 5, 1, 6), "대한민국 특허청\n상표 등록 현황", size=7.4)
    fill_cell(table.cell(0, 7), [("한국어 상표\n", {}), ("(해당 사항 O 표기)", {"color": "0000FF", "size": 6.8})], size=7.1)
    fill_cell(table.cell(0, 8), [("영어 상표\n", {}), ("(해당 사항 O 표기)", {"color": "0000FF", "size": 6.8})], size=7.1)
    fill_cell(table.cell(1, 7), "□미출원  □출원  □등록", size=6.5)
    fill_cell(table.cell(1, 8), "□미출원  □출원  □등록", size=6.5)

    fill_cell(merge_rect(table, 2, 0, 2, 2), "사업자등록번호", size=7.6)
    fill_cell(merge_rect(table, 2, 3, 2, 4), "")
    fill_cell(merge_rect(table, 2, 5, 2, 6), "법인등록번호", size=7.6)
    fill_cell(merge_rect(table, 2, 7, 2, 8), "")

    fill_cell(merge_rect(table, 3, 0, 3, 2), "대표자(성명)", size=7.6)
    fill_cell(merge_rect(table, 3, 3, 3, 4), "")
    fill_cell(merge_rect(table, 3, 5, 3, 6), "설립연월일", size=7.6)
    fill_cell(merge_rect(table, 3, 7, 3, 8), "        년      월      일", size=7.3)

    fill_cell(merge_rect(table, 4, 0, 6, 2), "담당자", size=8.0)
    contact_rows = [
        (4, "(성명)", "(직위)"),
        (5, "(전화)", "(핸드폰)"),
        (6, "(E-mail)", "(홈페이지)"),
    ]
    for r, left_label, right_label in contact_rows:
        fill_cell(table.cell(r, 3), left_label, size=7.3)
        fill_cell(merge_rect(table, r, 4, r, 5), "")
        fill_cell(table.cell(r, 6), right_label, size=7.1)
        fill_cell(merge_rect(table, r, 7, r, 8), "")

    fill_cell(merge_rect(table, 7, 0, 8, 0), "국내\n주소", size=7.8)
    fill_cell(merge_rect(table, 7, 1, 7, 2), "본  사", size=7.5)
    fill_cell(merge_rect(table, 7, 3, 7, 8), "□□□□□", align=WD_ALIGN_PARAGRAPH.LEFT, size=7.0)
    fill_cell(merge_rect(table, 8, 1, 8, 2), "공  장", size=7.5)
    fill_cell(merge_rect(table, 8, 3, 8, 8), "□□□□□", align=WD_ALIGN_PARAGRAPH.LEFT, size=7.0)

    fill_cell(merge_rect(table, 9, 0, 9, 2), "신청기업\n국내 매출액(`25)", size=7.0)
    fill_cell(merge_rect(table, 9, 3, 9, 4), [("KRW", {}), ("\n※재무제표 또는 부가세과세표준 증명원 기준", {"size": 5.6})], align=WD_ALIGN_PARAGRAPH.LEFT, size=7.0)
    fill_cell(merge_rect(table, 9, 5, 9, 6), "신청기업\n수출액(`25)", size=7.0)
    fill_cell(merge_rect(table, 9, 7, 9, 8), [("USD", {}), ("\n※한국무역통계진흥원에서 조회한 펫푸드 수출액 기준", {"size": 5.4})], align=WD_ALIGN_PARAGRAPH.LEFT, size=7.0)

    fill_cell(merge_rect(table, 10, 0, 10, 2), "국내외 인증서\n취득 현황", size=7.1)
    fill_cell(merge_rect(table, 10, 3, 10, 8), "美FDA(Certification, Approval), VEGAN, Gluten-free, NON-GMO, HALAL 등", align=WD_ALIGN_PARAGRAPH.LEFT, size=7.1, color="0000FF")

    # Product application section
    side = merge_rect(table, 11, 0, 20, 1)
    fill_cell(
        side,
        [
            ("입점\n신청\n주요\n상품\n\n", {"size": 8.0}),
            ("(2개\nSKU\n입점\n예정)", {"size": 8.0, "color": "FF0000"}),
        ],
        size=8.0,
    )

    fill_cell(table.cell(11, 2), "구분", size=7.5)
    fill_cell(merge_rect(table, 11, 3, 11, 8), [("제품(SKU) 1 :  ", {"bold": True}), ("제품명 기입", {"bold": True, "color": "0000FF"})], size=8.0)
    fill_cell(table.cell(12, 2), "상품\n품목", size=7.4)
    fill_cell(merge_rect(table, 12, 3, 12, 8), [
        ("- 중분류 : ", {}), ("상품 수출 HS CODE(품목번호 및 품명) 기입", {"color": "0000FF"}),
        ("\n- 소분류 : ", {}), ("기업 자율적으로 기입(1줄 이내) EX)딸기맛, 초코맛, 바나나맛", {"color": "0000FF"}),
    ], align=WD_ALIGN_PARAGRAPH.LEFT, size=7.0)
    fill_cell(table.cell(13, 2), "유통\n기한", size=7.4)
    fill_cell(merge_rect(table, 13, 3, 13, 8), [("- 유통기한(상온 보관 시) : ", {}), ("OOO일", {"color": "0000FF"})], align=WD_ALIGN_PARAGRAPH.LEFT, size=7.1)
    fill_cell(table.cell(14, 2), "상품\n특성", size=7.4)
    fill_cell(merge_rect(table, 14, 3, 14, 8), [
        ("- 국내 공급가(낱개) : ", {}), ("OOO원", {"color": "0000FF"}),
        (", 무게(낱개) : ", {}), ("OOOg", {"color": "0000FF"}),
        (", 부피(낱개) : ", {}), ("OOOml", {"color": "0000FF"}),
        (", MOQ : ", {}), ("OOO개", {"color": "0000FF"}),
        ("\n※ 입점 후 판매 활성화를 위해 입점 상품에 대한 국내 공급가는 국내 소비자 판매가의 약 50% 내외 기입 권장", {"size": 6.2}),
        ("\n- 기타 : ", {}), ("업체 자율적으로 작성(2줄 이내)", {"color": "0000FF"}),
    ], align=WD_ALIGN_PARAGRAPH.LEFT, size=7.0)
    fill_cell(table.cell(15, 2), "입점\n현황", size=7.4)
    fill_cell(merge_rect(table, 15, 3, 15, 8), [
        ("- 국내 : ", {}), ("매장명(온라인/오프라인 중 택1), 매장명(온라인/오프라인 중 택1), …", {"color": "0000FF"}),
        ("\n  ※ 인지도 높은 대형 유통매장 중심으로 3개 이내 작성", {"size": 6.2}),
        ("\n- 해외 : ", {}), ("매장명(국가명, 온라인/오프라인 중 택1), 매장명(국가명, 온라인/오프라인 중 택1), …", {"color": "0000FF"}),
        ("\n  ※ 인지도 높은 대형 유통매장 중심으로 3개 이내 작성", {"size": 6.2}),
    ], align=WD_ALIGN_PARAGRAPH.LEFT, size=6.7)

    fill_cell(table.cell(16, 2), "구분", size=7.5)
    fill_cell(merge_rect(table, 16, 3, 16, 8), [("제품(SKU) 2 :  ", {"bold": True}), ("제품명 기입", {"bold": True, "color": "0000FF"})], size=8.0)
    fill_cell(table.cell(17, 2), "상품\n품목", size=7.4)
    fill_cell(merge_rect(table, 17, 3, 17, 8), [
        ("- 중분류 : ", {}), ("상품 수출 HS CODE(품목번호 및 품명) 기입", {"color": "0000FF"}),
        ("\n- 소분류 : ", {}), ("기업 자율적으로 기입(1줄 이내) EX)딸기맛, 초코맛, 바나나맛", {"color": "0000FF"}),
    ], align=WD_ALIGN_PARAGRAPH.LEFT, size=7.0)
    fill_cell(table.cell(18, 2), "유통\n기한", size=7.4)
    fill_cell(merge_rect(table, 18, 3, 18, 8), [("- 유통기한(상온 보관 시) : ", {}), ("OOO일", {"color": "0000FF"})], align=WD_ALIGN_PARAGRAPH.LEFT, size=7.1)
    fill_cell(table.cell(19, 2), "상품\n특성", size=7.4)
    fill_cell(merge_rect(table, 19, 3, 19, 8), [
        ("- 국내 공급가(낱개) : ", {}), ("OOO원", {"color": "0000FF"}),
        (", 무게(낱개) : ", {}), ("OOOg", {"color": "0000FF"}),
        (", 부피(낱개) : ", {}), ("OOOml", {"color": "0000FF"}),
        (", MOQ : ", {}), ("OOO개", {"color": "0000FF"}),
        ("\n※ 입점 후 판매 활성화를 위해 입점 상품에 대한 국내 공급가는 국내 소비자 판매가의 약 50% 내외 기입 권장", {"size": 6.2}),
        ("\n- 기타 : ", {}), ("업체 자율적으로 작성(2줄 이내)", {"color": "0000FF"}),
    ], align=WD_ALIGN_PARAGRAPH.LEFT, size=7.0)
    fill_cell(table.cell(20, 2), "입점\n현황", size=7.4)
    fill_cell(merge_rect(table, 20, 3, 20, 8), [
        ("- 국내 : ", {}), ("매장명(온라인/오프라인 중 택1), 매장명(온라인/오프라인 중 택1), …", {"color": "0000FF"}),
        ("\n  ※ 인지도 높은 대형 유통매장 중심으로 3개 이내 작성", {"size": 6.2}),
        ("\n- 해외 : ", {}), ("매장명(국가명, 온라인/오프라인 중 택1), 매장명(국가명, 온라인/오프라인 중 택1), …", {"color": "0000FF"}),
        ("\n  ※ 인지도 높은 대형 유통매장 중심으로 3개 이내 작성", {"size": 6.2}),
    ], align=WD_ALIGN_PARAGRAPH.LEFT, size=6.7)

    set_repeat_table_header(table.rows[0])
    doc.core_properties.title = "일본 온라인 플랫폼 큐텐 입점지원 신청서"
    doc.core_properties.subject = "원본 HWPX 양식을 편집 가능한 DOCX로 변환"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_docx())
