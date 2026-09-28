from pathlib import Path
from decimal import Decimal
import json

from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r'C:\Users\Administrator\Documents\New project')
OUT = ROOT / 'outputs' / 'osp_sourcing_report_20260928'
OUT.mkdir(parents=True, exist_ok=True)
DOCX = OUT / 'OSP_동결건조_소싱_중간보고_20260928.docx'
FONT = '맑은 고딕'
doc = Document()
sec = doc.sections[0]
sec.page_width = Mm(210)
sec.page_height = Mm(297)
sec.top_margin = Mm(16)
sec.bottom_margin = Mm(15)
sec.left_margin = Mm(17)
sec.right_margin = Mm(17)
sec.footer_distance = Mm(7)
for grid in list(sec._sectPr.findall(qn('w:docGrid'))):
    sec._sectPr.remove(grid)
for style in doc.styles:
    if style.element.pPr is not None:
        for border in list(style.element.pPr.findall(qn('w:pBdr'))):
            style.element.pPr.remove(border)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = FONT
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = Pt(14)
    snap = OxmlElement('w:snapToGrid')
    snap.set(qn('w:val'), '0')
    style.element.get_or_add_pPr().append(snap)

doc.styles['Normal'].font.size = Pt(10)
doc.styles['Title'].font.size = Pt(19)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(5)
doc.styles['Title'].paragraph_format.line_spacing = Pt(25)
doc.styles['Heading 1'].font.size = Pt(11.5)
doc.styles['Heading 1'].font.bold = True
doc.styles['Heading 1'].paragraph_format.space_before = Pt(9)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(4)
doc.styles['Heading 1'].paragraph_format.keep_with_next = True
doc.styles['Heading 1'].paragraph_format.line_spacing = Pt(17)
for style in [doc.styles['Normal'], doc.styles['Title'], doc.styles['Heading 1']]:
    rpr = style.element.get_or_add_rPr()
    spacing = rpr.find(qn('w:spacing'))
    if spacing is None:
        spacing = OxmlElement('w:spacing')
        rpr.append(spacing)
    spacing.set(qn('w:val'), '0')
doc.core_properties.title = '동결건조 간식 소싱 중간보고'
doc.core_properties.author = 'OSP 최인호'
doc.core_properties.subject = '목표 판매가격과 공급사 협의 현황 및 샘플 검토'


def para(text, *, size=None, bold=False, after=4, before=0, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.keep_with_next = keep
    r = p.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    return p


def note(text):
    p = para(text, size=8.3, after=3)
    p.paragraph_format.line_spacing = Pt(11)
    return p


def heading(text):
    return doc.add_paragraph(text, style='Heading 1')


def table(headers, rows, widths, *, numeric=(), font_size=9.2):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), 'D9D9D9')
        borders.append(el)
    props.append(borders)
    for col, width in zip(t.columns, widths):
        col.width = Mm(width)
    header_mark = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(header_mark)
    for index, values in enumerate([headers] + rows):
        row = t.rows[0] if index == 0 else t.add_row()
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
        for j, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge, val in [('top', 62), ('bottom', 62), ('left', 85), ('right', 85)]:
                el = OxmlElement('w:' + edge)
                el.set(qn('w:w'), str(val))
                el.set(qn('w:type'), 'dxa')
                margins.append(el)
            tcpr.append(margins)
            if index == 0:
                fill = OxmlElement('w:shd')
                fill.set(qn('w:fill'), 'E6EBEF')
                tcpr.append(fill)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(12.5)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in numeric or index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            r.font.size = Pt(font_size)
            r.bold = index == 0
    return t


def hyperlink(p, text, url):
    rel = p.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), rel)
    r = OxmlElement('w:r')
    rp = OxmlElement('w:rPr')
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '16')
    rp.append(sz)
    r.append(rp)
    text_el = OxmlElement('w:t')
    text_el.text = text
    r.append(text_el)
    link.append(r)
    p._p.append(link)


doc.add_paragraph('동결건조 간식 소싱 중간보고', style='Title')
para('2026년 9월 28일  |  Brand영업팀 최인호  |  보고 대상 부장님', size=9, after=8)
para('네이처펫 대비 중량당 가격 경쟁력을 목표로 후난 페토와 라노바를 비교 중입니다. 인쇄포장과 국내 입고비용을 포함한 원가는 미확정이며, 샘플 평가와 동일 조건 재견적 후 품목별 진행 여부를 판단하고자 합니다.', after=6)

heading('1 개발 기준과 시장조사')
para('기존 기획 기준은 정상가 약 12,000원, 행사가 9,900원입니다. 최근 검토안은 표시가 11,900원에 10% 쿠폰을 적용한 10,710원이며, 최종 판매가격과 중량은 미확정입니다. 치킨 300g은 비교 기준이고, 전 품목 중량을 동일하게 고정하지 않습니다.')
table(['당시 비교 제품', '중량', '조사 가격', '검토 방향'], [
    ['네이처펫 치킨 큐브', '300g', '9,900원 쿠폰가', '주력 비교 품목'],
    ['네이처펫 치킨 텐더', '300g', '11,570원', '치킨 제형별 가격 참고'],
    ['네이처펫 북어 / 열빙어', '140g / 210g', '각 10,900원', '대구와 북어는 별도 원물'],
    ['추가 비교 메추리 난황', '200g', '9,290~9,580원', '당사는 닭 통난황 검토'],
], [54, 30, 42, 50], numeric=(1, 2))
note('9월 4일 및 7일 조사 기록이며 현재 판매가는 재확인 필요. 초기 정상가 10,900~11,900원 논의 후 약 12,000원으로 원가 역산. 쿠팡 마진은 33%에서 35%로 조정.')

heading('2 공급사 단가와 협의 현황')
note('기존 제품 견적 비교 / kg당 미달러 / 인쇄포장 완제품 최종 단가 아님')
table(['품목', '후난 페토', '라노바', '비교 시 확인사항'], [
    ['치킨', '10.90', '14.75', '큐브 규격 및 포장비'],
    ['대구', '36.60 / 26.00', '47.15', '후난 일반 / 저가 등급'],
    ['새우', '86.60', '124.30', '크기 및 가열 여부'],
    ['알배기 열빙어', '37.50', '46.50', '원물 크기 및 규격'],
    ['열빙어 암수 혼합 / 무란', '27.50 혼합', '19.75 무란', '혼합과 무란은 다른 사양'],
    ['난황', '8.00 막 포함', '8.37 통난황', '후난 통형 가격 적용 확인'],
], [44, 34, 34, 64], numeric=(1, 2))
note('후난 9월 17일 견적은 FOB 표기 및 200~500kg 구간. 라노바는 FOB 텐진 회신이며 가격 단위 재확인 필요. 후난 혼합 열빙어는 암수 각 50%, 라노바 무란은 수컷 100%로 단정 불가.')
para('후난 페토: OSP 지정 상해 창고까지 본품 운송비를 부담하겠다고 회신했습니다. 포장재 조달은 업체, 디자인은 OSP 담당으로 협의했습니다. 통난황 생산은 가능하나 생산수량 조건 및 최종 인쇄포장비는 확인이 필요합니다.', size=9.5)
note('후난 인쇄포장 제안: 품목별 약 2만개 기준 봉투 0.20~0.24달러/개 및 색상당 초기 인쇄비 105달러. 소량은 동일 크기 디지털 인쇄 가능하나 단가 미확정. 봉투값은 완제품 포장 총비용과 다릅니다.')
para('라노바: 벌크 최소주문량은 품목별 150kg입니다. 품목별 200kg을 비교 기준으로 인쇄 지퍼파우치 포함 FOB 텐진 견적을 요청했고 회신 대기 중입니다. 명태 견적은 아직 없습니다.', size=9.5)
note('라노바 비교 요청 중량: 치킨 300g, 대구 180g, 새우·열빙어 150g, 통난황 400g. 견적 비교용이며 최종 상품 중량이 아닙니다. 원더풀키친은 공유된 견적이 없어 수치 비교에서 제외했습니다.')

page2_title = doc.add_paragraph('수익성 검토와 진행 요청', style='Title')
page2_title.paragraph_format.page_break_before = True
para('기존 목표가격을 기준으로 허용 원가를 역산하고 샘플 및 물류 조건을 확정합니다', size=9, after=5)

heading('3 쿠폰과 광고비를 반영한 손익')
para('쿠폰과 행사 할인액은 OSP가 전액 부담합니다. 계산상 쿠팡 몫은 쿠폰 전 표시가의 35%로 고정하고, 광고비는 쿠폰 차감 후 OSP 매출의 5~10%로 가정했습니다. 실제 정산 및 광고 집행 기준은 별도 확인이 필요합니다.', size=9.5)
note('OSP 매출 = (표시가 × 65% - OSP 부담 할인액) ÷ 1.1. 아래 마진은 해당 매출 대비 원가·광고비 차감 후 잔여이익률이며 영업이익률은 아닙니다.')
cost = Decimal('3.76') * Decimal('1400')
price = Decimal('11900')
finance_rows = []
calc = []
for label, consumer in [('쿠폰 미적용', Decimal('11900')), ('상시 쿠폰 10%', Decimal('10710')), ('행사 최종가', Decimal('9900'))]:
    revenue = (price * Decimal('.65') - (price - consumer)) / Decimal('1.1')
    m5 = (revenue * Decimal('.95') - cost) / revenue * 100
    m10 = (revenue * Decimal('.90') - cost) / revenue * 100
    finance_rows.append([label, f'{consumer:,.0f}원', f'{revenue:,.0f}원', f'{m5:.1f}%', f'{m10:.1f}%'])
    calc.append(dict(scenario=label, consumer=float(consumer), revenue=float(revenue), margin_ad5=float(m5), margin_ad10=float(m10)))
table(['표시가 11,900원', '소비자 결제', 'OSP 매출', '광고 5%\n차감 후', '광고 10%\n차감 후'], finance_rows, [43, 35, 34, 32, 32], numeric=(1, 2, 3, 4))
note('참고 원가: 후난 치킨 300g 기존 견적 3.76달러 × 가정 환율 1,400원 = 5,264원. 무지봉투·스티커 기준이며 추가 수입·국내 물류비 및 인쇄포장 변경분은 미반영. 실제 마진은 최종 원가로 재산정합니다.')
para('상시 쿠폰과 광고비 10% 적용 시 잔여 마진은 추가 비용 전에도 1.5%입니다. 기존 20% 목표를 광고비 차감 후에도 유지하려면 총원가는 4,165원 이하여야 합니다. 단가 인하 없이 가격만 높이면 네이처펫 대비 가격 우위가 약해집니다.', size=9.5, bold=True)

heading('4 샘플 평가와 특송비')
para('평가 목적은 품질, 제형 및 기호성 확인입니다. 원물 크기·균일성, 파손·가루 발생, 식감·냄새와 급여 반응을 확인하며, 품목별 필요량은 시험 규모에 맞춰 결정합니다.', size=9.5)
para('후난은 샘플 제품 무상, 특송비 OSP 부담 조건을 제시했습니다. 한국 직송 기준 특송료와 발송일을 확인할 예정입니다. 라노바의 샘플비 및 특송비 부담 조건은 미확정입니다. 본품의 상해 창고 무상운송 조건과 샘플 특송은 별도입니다.', size=9.5)

heading('5 물류 조건 비교')
table(['비교안', '공급사 인도 기준', 'OSP 측 추가비용 비교 구간'], [
    ['후난 FCA 상해 지정 창고', '지정 창고 차량 도착\n하역 준비 상태 인도', '창고 하역·항만 반입·선적 이후\n동일한 국내 창고 입고까지'],
    ['후난 FOB 상해 / 라노바 FOB 텐진', '각 선적항 본선 적재 완료', '본선 적재 이후\n동일한 국내 창고 입고까지'],
], [58, 50, 68], font_size=9)
note('FCA는 운송인 인도, FOB는 본선 인도 조건입니다. 두 조건 모두 수출통관은 공급사 책임이 원칙이며 견적서에 인도 장소와 포함 비용을 명시해야 합니다.')
para('후난의 생산·포장은 후난 공장에서 진행되므로 상해 인건비가 절감 근거는 아닙니다. FCA의 유불리는 지정 창고 이후 포워더 비용과 FOB 가산비용을 비교해야 판단할 수 있습니다. 항구를 억지로 통일하지 않고 국내 창고 입고원가로 비교합니다.', size=9.5)

heading('6 내부 결정과 후속 요청')
para('① 샘플 우선 품목·품목별 필요량 결정 및 특송 견적 수령 후 비용 승인\n② 평시와 행사 각각의 최소 허용 마진 및 광고비 집행 기준 결정\n③ 양사 인쇄포장 완제품 재견적과 포워더 운임 확보 후 최종 중량·발주량 검토', size=9.5, after=4)
note('자료: 9월 4일·7일 시장조사 기록, 후난 9월 17일 견적 및 후속 메일, 라노바 최근 FOB 텐진 회신. 기존 가격·사양 차이는 최종 재견적에서 정합화합니다.')
p = note('물류 조건 참고: ')
hyperlink(p, 'ICC Academy의 FCA 및 FOB 해설', 'https://academy.iccwbo.org/incoterms/article/incoterms-2020-fca-or-fob/')

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run('OSP  |  ')
run.font.name = FONT
run.font.size = Pt(8)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

doc.save(DOCX)
(OUT / 'calculation_check.json').write_text(json.dumps({'cost_assumption': float(cost), 'rows': calc, 'regular_cost_ceiling_ad10_margin20': 4165}, ensure_ascii=False, indent=2), encoding='utf-8')
print(DOCX)
